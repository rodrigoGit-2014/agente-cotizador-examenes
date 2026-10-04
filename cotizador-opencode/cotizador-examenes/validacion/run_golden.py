"""Harness del golden set (paso `11`, agente `validator`).

Ejecuta cada caso, captura el paso a paso y compara contra `esperado`. Escribe
`resultados_{fecha}.xlsx` con las hojas Detalle, Resumen y Fallos.

Uso:
    python validacion/run_golden.py [--limit N] [--solo-id G01] [--salida ruta.xlsx]
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from langchain_core.messages import AIMessage  # noqa: E402

import main  # noqa: E402
from services import privacidad  # noqa: E402

AQUI = Path(__file__).resolve().parent


def _herramientas_usadas(mensajes: list) -> list[str]:
    nombres = []
    for m in mensajes:
        for tc in getattr(m, "tool_calls", []) or []:
            nombres.append(tc.get("name", ""))
    return nombres


def ejecutar_caso(caso: dict) -> dict:
    thread = {"configurable": {"thread_id": f"golden-{caso['id']}-{time.time()}"}}
    nodos: list[str] = []
    error = ""
    estado: dict = {}
    t0 = time.time()
    try:
        for turno in caso["turnos"]:
            entrada = privacidad.preparar_entrada(turno)
            for chunk in main.app.stream(
                {"messages": [entrada], "escenario_id": caso["escenario"]},
                thread,
                stream_mode="updates",
            ):
                for nodo in chunk:
                    nodos.append(nodo)
        estado = main.app.get_state(thread).values or {}
    except Exception as exc:  # una pregunta no aborta el resto
        error = f"{type(exc).__name__}: {exc}"
    return {"estado": estado, "nodos": nodos, "error": error, "latencia": round(time.time() - t0, 2)}


def _comparar(caso: dict, res: dict) -> list[str]:
    exp = caso.get("esperado", {})
    estado = res["estado"]
    sol = estado.get("solicitud")
    rep = estado.get("reporte")
    usadas = set(_herramientas_usadas(estado.get("messages", [])))
    fallos: list[str] = []

    def check(ok: bool, etiqueta: str, obtenido) -> None:
        if not ok:
            fallos.append(f"{etiqueta} (obtenido: {obtenido})")

    if "ruta" in exp:
        check(estado.get("ruta") == exp["ruta"], "ruta", estado.get("ruta"))
    if "solicitud" in exp and sol is not None:
        for clave, valor in exp["solicitud"].items():
            actual = sol.prevision.value if clave == "prevision" else getattr(sol, clave, None)
            check(str(actual) == str(valor), f"solicitud.{clave}", actual)
    if "n_comparables" in exp:
        check(rep is not None and len(rep.comparables) == exp["n_comparables"],
              "n_comparables", len(rep.comparables) if rep else None)
    if "motivo_parada" in exp:
        check(estado.get("motivo_parada") == exp["motivo_parada"], "motivo_parada", estado.get("motivo_parada"))
    if rep is not None and "comparables" in exp:
        esperados = {(c["centro_id"], c.get("codigo_examen")) for c in exp["comparables"]}
        reales = {(c.centro_id, c.codigo_examen) for c in rep.comparables}
        check(esperados <= reales, "comparables", sorted(reales))
    if rep is not None and "descartados_contiene" in exp:
        reales = {(d.centro_id, d.motivo) for d in rep.descartados}
        for esperado in exp["descartados_contiene"]:
            check((esperado["centro_id"], esperado["motivo"]) in reales, "descartados", sorted(reales))
    if rep is not None and "dudosos" in exp:
        reales = {d.centro_id for d in rep.dudosos}
        check(set(exp["dudosos"]) <= reales, "dudosos", sorted(reales))
    if "usa" in exp:
        for nombre in exp["usa"]:
            check(nombre in usadas, f"usa {nombre}", sorted(usadas))
    if "no_debe_usar" in exp:
        for nombre in exp["no_debe_usar"]:
            check(nombre not in usadas, f"no_debe_usar {nombre}", sorted(usadas))
    if "horas" in exp and rep is not None:
        por_id = {c.centro_id: c for c in rep.comparables}
        for centro_id, esperado in exp["horas"].items():
            cot = por_id.get(centro_id)
            check(cot is not None and cot.proxima_hora.estado == esperado,
                  f"hora {centro_id}", cot.proxima_hora.estado if cot else None)
    if exp.get("reserva_no_aceptada") and rep is not None:
        check(all(not c.reserva.aceptada for c in rep.comparables), "reserva_no_aceptada", "alguna aceptada")
    if "respuesta_contiene_alguno" in exp:
        final = estado.get("messages", [])
        texto = (getattr(final[-1], "content", "") if final else "").lower()
        check(any(f.lower() in texto for f in exp["respuesta_contiene_alguno"]),
              "respuesta_contiene_alguno", texto[:80])
    if "transcripcion_no_contiene" in exp and rep is not None:
        todo = " ".join(
            t.texto for c in rep.comparables for t in c.transcripcion
        )
        for prohibido in exp["transcripcion_no_contiene"]:
            check(prohibido not in todo, f"transcripcion_no_contiene {prohibido}", "presente")

    return fallos


def main_cli() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--solo-id", default=None)
    parser.add_argument("--salida", type=Path, default=None)
    args = parser.parse_args()

    casos = json.loads((AQUI / "golden_set.json").read_text(encoding="utf-8"))
    if args.solo_id:
        casos = [c for c in casos if c["id"] == args.solo_id]
    if args.limit:
        casos = casos[: args.limit]

    filas = []
    for i, caso in enumerate(casos, 1):
        print(f"[{i}/{len(casos)}] {caso['id']} ...")
        res = ejecutar_caso(caso)
        fallos = _comparar(caso, res) if not res["error"] else [res["error"]]
        filas.append({
            "id": caso["id"], "situacion": caso["situacion"], "escenario": caso["escenario"],
            "ok": not fallos, "fallos": "; ".join(fallos),
            "traza": " -> ".join(res["nodos"]), "latencia": res["latencia"],
            "herramientas": ", ".join(sorted(set(_herramientas_usadas(res["estado"].get("messages", []))))),
        })

    salida = args.salida or (AQUI / f"resultados_{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    _escribir(salida, filas)
    aprobados = sum(1 for f in filas if f["ok"])
    print(f"\nAprobados: {aprobados}/{len(filas)}. Informe: {salida}")


def _escribir(ruta: Path, filas: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    detalle = wb.active
    detalle.title = "Detalle"
    cabecera = ["id", "situacion", "escenario", "ok", "fallos", "traza", "herramientas", "latencia_s"]
    detalle.append(cabecera)
    for f in filas:
        detalle.append([f["id"], f["situacion"], f["escenario"], f["ok"], f["fallos"],
                        f["traza"], f["herramientas"], f["latencia"]])
    for celda in detalle[1]:
        celda.font = Font(bold=True)
    detalle.freeze_panes = "A2"

    resumen = wb.create_sheet("Resumen")
    aprobados = sum(1 for f in filas if f["ok"])
    resumen.append(["casos", len(filas)])
    resumen.append(["aprobados", aprobados])
    resumen.append(["fallidos", len(filas) - aprobados])
    latencias = [f["latencia"] for f in filas if f["latencia"] is not None]
    if latencias:
        resumen.append(["latencia_media", round(sum(latencias) / len(latencias), 2)])
        resumen.append(["latencia_max", max(latencias)])

    fallos = wb.create_sheet("Fallos")
    fallos.append(["id", "situacion", "fallos"])
    for f in filas:
        if not f["ok"]:
            fallos.append([f["id"], f["situacion"], f["fallos"]])
    for celda in fallos[1]:
        celda.font = Font(bold=True)
    wb.save(ruta)


if __name__ == "__main__":
    main_cli()
